"""Generate test fixtures with known, planted metadata.

Every fixture contains the marker strings below in its metadata so tests can
check that they are gone (or pseudonymised) after cleaning. Nothing here is a
binary of unknown origin: all files are produced from scratch with PIL,
pikepdf, python-docx, openpyxl, python-pptx, ExifTool and FFmpeg.

Usage: python3 -m tests.fixtures.build_fixtures <output_dir>
"""

from __future__ import annotations

import io
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

AUTHOR = "Fer Autor"
SERIAL = "SN-998877"
PRIVATE_PATH = "C:\\Users\\fer\\Documentos\\secreto"
COMMENT_TEXT = "comentario privado"
BODY_TEXT = "Texto visible del documento"

MARKERS = (AUTHOR, SERIAL)


def run(*args: str) -> None:
    subprocess.run(list(args), check=True, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)


def exif_tag(path: Path, *tags: str) -> None:
    run("exiftool", "-q", "-q", "-m", "-overwrite_original", *tags, str(path))


COMMON_TAGS = (f"-Artist={AUTHOR}", f"-XMP-dc:Creator={AUTHOR}", f"-SerialNumber={SERIAL}",
               "-Make=MarcaCam", "-Model=ModeloX", "-DateTimeOriginal=2023:05:06 07:08:09",
               "-GPSLatitude=33.4489", "-GPSLatitudeRef=S", "-GPSLongitude=70.6693", "-GPSLongitudeRef=W")


# ---------------------------------------------------------------- images
def _gradient(w: int, h: int, mode: str = "RGB"):
    from PIL import Image
    im = Image.new(mode, (w, h))
    px = im.load()
    for y in range(h):
        for x in range(w):
            v = ((x * 255) // max(w - 1, 1), (y * 255) // max(h - 1, 1), ((x + y) * 7) % 256)
            px[x, y] = v + ((x * 3 + 60) % 256,) if mode == "RGBA" else v
    return im


def _icc() -> bytes | None:
    try:
        from PIL import ImageCms
        return ImageCms.ImageCmsProfile(ImageCms.createProfile("sRGB")).tobytes()
    except Exception:
        return None


def jpeg(out: Path) -> Path:
    p = out / "photo.jpg"
    icc = _icc()
    kwargs = {"quality": 90}
    if icc:
        kwargs["icc_profile"] = icc
    _gradient(96, 64).save(p, **kwargs)
    thumb = out / "_thumb.jpg"
    _gradient(16, 12).save(thumb, quality=70)
    exif_tag(p, *COMMON_TAGS, "-Orientation#=6", f"-IPTC:By-line={AUTHOR}", f"-Comment={AUTHOR}",
             f"-ThumbnailImage<={thumb}")
    thumb.unlink()
    return p


def png(out: Path) -> Path:
    from PIL.PngImagePlugin import PngInfo
    p = out / "image.png"
    info = PngInfo()
    info.add_text("Author", AUTHOR)
    info.add_text("Comment", f"serial {SERIAL}")
    _gradient(40, 30, "RGBA").save(p, pnginfo=info)
    exif_tag(p, f"-XMP-dc:Creator={AUTHOR}", "-XMP-exif:GPSLatitude=33.4489 S")
    return p


def webp(out: Path) -> Path:
    from PIL import Image
    p = out / "image.webp"
    exif = Image.Exif()
    exif[0x013B] = AUTHOR      # Artist
    exif[0xA431] = SERIAL      # BodySerialNumber
    _gradient(40, 30).save(p, lossless=True, exif=exif.tobytes())
    exif_tag(p, f"-XMP-dc:Creator={AUTHOR}")
    return p


def tiff(out: Path) -> Path:
    p = out / "scan.tiff"
    _gradient(40, 30).save(p)
    exif_tag(p, *COMMON_TAGS)
    return p


def gif(out: Path) -> Path:
    from PIL import Image
    p = out / "anim.gif"
    frames = [Image.new("RGB", (20, 20), c).convert("P") for c in ((255, 0, 0), (0, 255, 0), (0, 0, 255))]
    frames[0].save(p, save_all=True, append_images=frames[1:], duration=100, loop=0, comment=AUTHOR.encode())
    exif_tag(p, f"-XMP-dc:Creator={AUTHOR}")
    return p


def heic(out: Path) -> Path | None:
    if not shutil.which("heif-enc"):
        return None
    src = out / "_heic_src.png"
    _gradient(64, 64).save(src)
    p = out / "photo.heic"
    try:
        run("heif-enc", "-q", "60", "-o", str(p), str(src))
    except (subprocess.CalledProcessError, OSError):
        return None
    finally:
        src.unlink(missing_ok=True)
    exif_tag(p, *COMMON_TAGS)
    return p


# ---------------------------------------------------------------- pdf
def _pdf(path: Path, pages: int = 2, encrypt: bool = False) -> Path:
    import pikepdf
    pdf = pikepdf.new()
    font = pdf.make_indirect(pikepdf.Dictionary(Type=pikepdf.Name.Font, Subtype=pikepdf.Name.Type1,
                                                BaseFont=pikepdf.Name.Helvetica))
    for i in range(pages):
        pdf.add_blank_page(page_size=(612, 792))
        pg = pdf.pages[-1]
        pg.Resources = pikepdf.Dictionary(Font=pikepdf.Dictionary(F1=font))
        pg.Contents = pdf.make_stream(f"BT /F1 20 Tf 72 700 Td ({BODY_TEXT} {i + 1}) Tj ET".encode())
    note = pdf.make_indirect(pikepdf.Dictionary(
        Type=pikepdf.Name.Annot, Subtype=pikepdf.Name.Text, Rect=[100, 100, 120, 120],
        Contents=pikepdf.String(COMMENT_TEXT), T=pikepdf.String(AUTHOR), M=pikepdf.String("D:20230506070809")))
    pdf.pages[0].Annots = pdf.make_indirect(pikepdf.Array([note]))
    pdf.docinfo["/Author"] = AUTHOR
    pdf.docinfo["/Creator"] = f"Writer {SERIAL}"
    pdf.docinfo["/Title"] = "Documento de prueba"
    with pdf.open_metadata(set_pikepdf_as_editor=False) as m:
        m["dc:creator"] = [AUTHOR]
        m["xmp:CreatorTool"] = f"Tool {SERIAL}"
    attachment = pikepdf.AttachedFileSpec(pdf, b"adjunto", filename="adjunto.txt",
                                          creation_date="D:20230506070809", mod_date="D:20230506070809")
    pdf.attachments["adjunto.txt"] = attachment
    js = pdf.make_indirect(pikepdf.Dictionary(S=pikepdf.Name.JavaScript, JS=pikepdf.String("app.alert(1);")))
    pdf.Root.Names = pdf.Root.get("/Names", pikepdf.Dictionary())
    pdf.Root.Names.JavaScript = pikepdf.Dictionary(Names=pikepdf.Array([pikepdf.String("x"), js]))
    kwargs = {}
    if encrypt:
        kwargs["encryption"] = pikepdf.Encryption(owner="owner-pass", user="user-pass")
    pdf.save(path, **kwargs)
    return path


def pdf(out: Path) -> Path:
    return _pdf(out / "document.pdf")


def pdf_encrypted(out: Path) -> Path:
    return _pdf(out / "encrypted.pdf", encrypt=True)


# ---------------------------------------------------------------- office
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _patch_zip(path: Path, edits: dict[str, bytes | None], add_ct: list[str] = ()) -> None:
    """Rewrite a zip replacing/adding members (None deletes) and appending
    <Override> entries to [Content_Types].xml."""
    with zipfile.ZipFile(path) as zin:
        members = {i.filename: zin.read(i.filename) for i in zin.infolist()}
    for k, v in edits.items():
        if v is None:
            members.pop(k, None)
        else:
            members[k] = v
    if add_ct:
        ct = members["[Content_Types].xml"]
        members["[Content_Types].xml"] = ct.replace(b"</Types>", "".join(add_ct).encode() + b"</Types>")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name in ["[Content_Types].xml"] + [n for n in members if n != "[Content_Types].xml"]:
            zout.writestr(name, members[name])


def _add_rel(rels: bytes, rid: str, rtype: str, target: str, external: bool = False) -> bytes:
    mode = ' TargetMode="External"' if external else ""
    return rels.replace(b"</Relationships>",
                        f'<Relationship Id="{rid}" Type="{rtype}" Target="{target}"{mode}/></Relationships>'.encode())


CUSTOM_XML = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
              '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
              'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
              '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="Cliente">'
              f'<vt:lpwstr>{AUTHOR}</vt:lpwstr></property></Properties>').encode()
CUSTOM_CT = ('<Override PartName="/docProps/custom.xml" '
             'ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml"/>')
CUSTOM_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties"


def _with_custom_props(path: Path) -> None:
    with zipfile.ZipFile(path) as zf:
        rels = zf.read("_rels/.rels")
    _patch_zip(path, {"docProps/custom.xml": CUSTOM_XML,
                      "_rels/.rels": _add_rel(rels, "rIdCustom1", CUSTOM_REL, "docProps/custom.xml")}, [CUSTOM_CT])


def docx(out: Path, image: Path | None = None) -> Path:
    import docx as pydocx
    p = out / "document.docx"
    d = pydocx.Document()
    d.add_paragraph(BODY_TEXT)
    d.add_paragraph("Elemento de lista", style="List Bullet")
    if image is not None:
        d.add_picture(str(image))
    cp = d.core_properties
    cp.author = AUTHOR
    cp.last_modified_by = AUTHOR
    cp.comments = f"serial {SERIAL}"
    cp.title = "Informe"
    d.save(p)
    with zipfile.ZipFile(p) as zf:
        doc = zf.read("word/document.xml").decode()
        doc_rels = zf.read("word/_rels/document.xml.rels")
        settings_xml = zf.read("word/settings.xml").decode()
        names = zf.namelist()
    extra = (f'<w:p w:rsidR="00A1B2C3" w:rsidRDefault="00A1B2C3"><w:commentRangeStart w:id="0"/>'
             f'<w:r><w:t>Parrafo comentado</w:t></w:r><w:commentRangeEnd w:id="0"/>'
             f'<w:r><w:commentReference w:id="0"/></w:r>'
             f'<w:ins w:id="1" w:author="{AUTHOR}" w:date="2023-05-06T07:08:09Z">'
             f'<w:r><w:t>texto insertado</w:t></w:r></w:ins></w:p>')
    doc = doc.replace("<w:sectPr", extra + "<w:sectPr", 1)
    comments = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:comments xmlns:w="{W_NS}">'
                f'<w:comment w:id="0" w:author="{AUTHOR}" w:date="2023-05-06T07:08:09Z" w:initials="FA">'
                f'<w:p><w:r><w:t>{COMMENT_TEXT}</w:t></w:r></w:p></w:comment></w:comments>').encode()
    doc_rels = _add_rel(doc_rels, "rIdCmt1",
                        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments",
                        "comments.xml")
    settings_xml = settings_xml.replace("<w:defaultTabStop", '<w:attachedTemplate r:id="rIdTpl1"/><w:defaultTabStop', 1)
    if "xmlns:r=" not in settings_xml.split(">", 2)[1]:
        settings_xml = settings_xml.replace(
            "<w:settings ", '<w:settings xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" ', 1)
    settings_rels = _add_rel(
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"></Relationships>',
        "rIdTpl1", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/attachedTemplate",
        "file:///" + PRIVATE_PATH.replace("\\", "/") + "/plantilla.dotx", external=True)
    edits = {"word/document.xml": doc.encode(), "word/comments.xml": comments,
             "word/_rels/document.xml.rels": doc_rels, "word/settings.xml": settings_xml.encode(),
             "word/_rels/settings.xml.rels": settings_rels}
    assert "word/_rels/settings.xml.rels" not in names
    _patch_zip(p, edits, ['<Override PartName="/word/comments.xml" ContentType="application/'
                          'vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"/>'])
    _with_custom_props(p)
    return p


def xlsx(out: Path) -> Path:
    import openpyxl
    from openpyxl.comments import Comment
    p = out / "sheet.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Datos"
    ws["A1"], ws["A2"], ws["A3"] = 2, 3, "=A1*A2"
    ws["B1"] = BODY_TEXT
    ws["C1"] = "=SUM(A1:A3)"
    ws["B2"].comment = Comment(COMMENT_TEXT, AUTHOR)
    hidden = wb.create_sheet("Oculta")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "dato oculto"
    wb.properties.creator = AUTHOR
    wb.properties.lastModifiedBy = AUTHOR
    wb.properties.keywords = SERIAL
    wb.save(p)
    with zipfile.ZipFile(p) as zf:
        wbxml = zf.read("xl/workbook.xml").decode()
    abspath = ('<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
               '<mc:Choice Requires="x15"><x15ac:absPath xmlns:x15ac="http://schemas.microsoft.com/office/'
               f'spreadsheetml/2010/11/ac" url="{PRIVATE_PATH}\\"/></mc:Choice></mc:AlternateContent>')
    wbxml = re.sub(r"(<workbookPr\b[^>]*/>)", lambda m: m.group(1) + abspath, wbxml, count=1)
    assert "absPath" in wbxml
    _patch_zip(p, {"xl/workbook.xml": wbxml.encode()})
    _with_custom_props(p)
    return p


def pptx(out: Path) -> Path:
    from pptx import Presentation
    p = out / "slides.pptx"
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = BODY_TEXT
    s1.placeholders[1].text = "Punto uno"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Diapositiva oculta"
    prs.core_properties.author = AUTHOR
    prs.core_properties.last_modified_by = AUTHOR
    prs.core_properties.comments = SERIAL
    prs.save(p)
    with zipfile.ZipFile(p) as zf:
        rels = zf.read("ppt/_rels/presentation.xml.rels")
        slide2 = zf.read("ppt/slides/slide2.xml").decode()
    authors = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
               '<p:cmAuthorLst xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
               f'<p:cmAuthor id="0" name="{AUTHOR}" initials="FA" lastIdx="1" clrIdx="0"/></p:cmAuthorLst>').encode()
    rels = _add_rel(rels, "rIdAuth1",
                    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/commentAuthors",
                    "commentAuthors.xml")
    slide2 = re.sub(r"<p:sld\b", '<p:sld show="0"', slide2, count=1)
    _patch_zip(p, {"ppt/commentAuthors.xml": authors, "ppt/_rels/presentation.xml.rels": rels,
                   "ppt/slides/slide2.xml": slide2.encode()},
               ['<Override PartName="/ppt/commentAuthors.xml" ContentType="application/'
                'vnd.openxmlformats-officedocument.presentationml.commentAuthors+xml"/>'])
    return p


# ---------------------------------------------------------------- opendocument
ODF_NS = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
          'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
          'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
          'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" '
          'xmlns:presentation="urn:oasis:names:tc:opendocument:xmlns:presentation:1.0" '
          'xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" '
          'xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0" '
          'xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0" '
          'xmlns:dc="http://purl.org/dc/elements/1.1/" '
          'xmlns:meta="urn:oasis:names:tc:opendocument:xmlns:meta:1.0" '
          'xmlns:config="urn:oasis:names:tc:opendocument:xmlns:config:1.0" '
          'xmlns:officeooo="http://openoffice.org/2009/office"')
ANNOTATION = (f'<office:annotation><dc:creator>{AUTHOR}</dc:creator><dc:date>2023-05-06T07:08:09</dc:date>'
              f'<text:p>{COMMENT_TEXT}</text:p></office:annotation>')
ODF_BODIES = {
    "odt": ("application/vnd.oasis.opendocument.text",
            f'<office:text><text:p text:style-name="P1">{BODY_TEXT}{ANNOTATION}</text:p></office:text>'),
    "ods": ("application/vnd.oasis.opendocument.spreadsheet",
            '<office:spreadsheet><table:table table:name="Hoja1"><table:table-row>'
            f'<table:table-cell office:value-type="string"><text:p>{BODY_TEXT}</text:p>{ANNOTATION}</table:table-cell>'
            '<table:table-cell table:formula="of:=1+2" office:value-type="float" office:value="3"><text:p>3</text:p>'
            '</table:table-cell></table:table-row></table:table></office:spreadsheet>'),
    "odp": ("application/vnd.oasis.opendocument.presentation",
            '<office:presentation><draw:page draw:name="p1"><draw:frame svg:width="10cm" svg:height="2cm" '
            f'svg:x="1cm" svg:y="1cm"><draw:text-box><text:p>{BODY_TEXT}</text:p></draw:text-box></draw:frame>'
            f'{ANNOTATION}</draw:page></office:presentation>'),
}


def odf(out: Path, kind: str) -> Path:
    from PIL import Image
    mime, body = ODF_BODIES[kind]
    p = out / f"document.{kind}"
    content = (f'<?xml version="1.0" encoding="UTF-8"?><office:document-content {ODF_NS} office:version="1.3">'
               '<office:automatic-styles><style:style style:name="P1" style:family="paragraph">'
               '<style:text-properties officeooo:rsid="0012ab34"/></style:style></office:automatic-styles>'
               f'<office:body>{body}</office:body></office:document-content>')
    meta = (f'<?xml version="1.0" encoding="UTF-8"?><office:document-meta {ODF_NS} office:version="1.3"><office:meta>'
            f'<meta:initial-creator>{AUTHOR}</meta:initial-creator><dc:creator>{AUTHOR}</dc:creator>'
            f'<meta:generator>Suite/{SERIAL}</meta:generator><dc:date>2023-05-06T07:08:09</dc:date>'
            '<meta:editing-duration>PT1H</meta:editing-duration></office:meta></office:document-meta>')
    styles = f'<?xml version="1.0" encoding="UTF-8"?><office:document-styles {ODF_NS} office:version="1.3"/>'
    settings = (f'<?xml version="1.0" encoding="UTF-8"?><office:document-settings {ODF_NS} office:version="1.3">'
                '<office:settings><config:config-item-set config:name="ooo:configuration-settings">'
                f'<config:config-item config:name="PrinterName" config:type="string">Impresora de {AUTHOR}'
                '</config:config-item></config:config-item-set></office:settings></office:document-settings>')
    thumb = io.BytesIO()
    Image.new("RGB", (16, 16), (200, 200, 200)).save(thumb, "PNG")
    manifest = ('<?xml version="1.0" encoding="UTF-8"?><manifest:manifest '
                'xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.3">'
                f'<manifest:file-entry manifest:full-path="/" manifest:media-type="{mime}"/>'
                '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
                '<manifest:file-entry manifest:full-path="styles.xml" manifest:media-type="text/xml"/>'
                '<manifest:file-entry manifest:full-path="meta.xml" manifest:media-type="text/xml"/>'
                '<manifest:file-entry manifest:full-path="settings.xml" manifest:media-type="text/xml"/>'
                '<manifest:file-entry manifest:full-path="Thumbnails/thumbnail.png" manifest:media-type="image/png"/>'
                '</manifest:manifest>')
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(zipfile.ZipInfo("mimetype"), mime, compress_type=zipfile.ZIP_STORED)
        zf.writestr("content.xml", content)
        zf.writestr("styles.xml", styles)
        zf.writestr("meta.xml", meta)
        zf.writestr("settings.xml", settings)
        zf.writestr("Thumbnails/thumbnail.png", thumb.getvalue())
        zf.writestr("META-INF/manifest.xml", manifest)
    return p


# ---------------------------------------------------------------- audio / video
META = ["-metadata", f"artist={AUTHOR}", "-metadata", f"comment={SERIAL}", "-metadata", "title=Grabacion",
        "-metadata", "date=2023-05-06"]
FF = ["ffmpeg", "-nostdin", "-hide_banner", "-loglevel", "error", "-y"]


def audio(out: Path, kind: str, cover: Path | None) -> Path:
    p = out / f"audio.{kind}"
    src = ["-f", "lavfi", "-i", "sine=frequency=440:duration=2"]
    codec = {"mp3": ["-c:a", "libmp3lame", "-b:a", "64k"], "wav": ["-c:a", "pcm_s16le"],
             "flac": ["-c:a", "flac"], "m4a": ["-c:a", "aac", "-b:a", "64k"], "aac": ["-c:a", "aac", "-b:a", "64k"],
             "ogg": ["-c:a", "libvorbis", "-q:a", "2"]}[kind]
    if cover is not None and kind in ("mp3", "m4a", "flac"):
        args = FF + src + ["-i", str(cover), "-map", "0:a", "-map", "1:v", "-c:v", "copy" if kind != "m4a" else "mjpeg",
                           "-disposition:v", "attached_pic"] + codec + META
        if kind == "mp3":
            args += ["-id3v2_version", "3"]
    else:
        args = FF + src + codec + META
    if kind == "aac":
        args += ["-f", "adts", "-write_id3v2", "1"]
    run(*args, str(p))
    return p


def video(out: Path, kind: str) -> Path:
    p = out / f"video.{kind}"
    src = ["-f", "lavfi", "-i", "testsrc=size=160x120:rate=15:duration=2",
           "-f", "lavfi", "-i", "sine=frequency=330:duration=2"]
    codec = {"mp4": ["-c:v", "libx264", "-preset", "ultrafast", "-c:a", "aac"],
             "mov": ["-c:v", "libx264", "-preset", "ultrafast", "-c:a", "aac"],
             "mkv": ["-c:v", "libx264", "-preset", "ultrafast", "-c:a", "libvorbis"],
             "webm": ["-c:v", "libvpx", "-b:v", "200k", "-c:a", "libopus"],
             "avi": ["-c:v", "mpeg4", "-c:a", "libmp3lame"]}[kind]
    extra = ["-metadata", "location=-33.4489-070.6693/", "-metadata", "creation_time=2023-05-06T07:08:09Z"]
    if kind == "mkv":
        srt = out / "_subs.srt"
        srt.write_text("1\n00:00:00,000 --> 00:00:01,000\nHola\n", encoding="utf-8")
        run(*FF, *src, "-i", str(srt), "-map", "0", "-map", "1", "-map", "2", *codec, "-c:s", "srt",
            *META, *extra, str(p))
        srt.unlink()
    else:
        run(*FF, *src, *codec, *META, *extra, str(p))
    return p


def build_all(out: Path) -> dict[str, Path]:
    out.mkdir(parents=True, exist_ok=True)
    f: dict[str, Path] = {}
    f["jpeg"] = jpeg(out)
    f["png"] = png(out)
    f["webp"] = webp(out)
    f["tiff"] = tiff(out)
    f["gif"] = gif(out)
    h = heic(out)
    if h:
        f["heic"] = h
    f["pdf"] = pdf(out)
    f["pdf_encrypted"] = pdf_encrypted(out)
    f["docx"] = docx(out, image=f["jpeg"])
    f["xlsx"] = xlsx(out)
    f["pptx"] = pptx(out)
    for kind in ("odt", "ods", "odp"):
        f[kind] = odf(out, kind)
    cover = out / "_cover.jpg"
    _gradient(32, 32).save(cover, quality=80)
    exif_tag(cover, f"-Artist={AUTHOR}")
    for kind in ("mp3", "wav", "flac", "m4a", "aac", "ogg"):
        f[kind] = audio(out, kind, cover)
    cover.unlink()
    for kind in ("mp4", "mov", "mkv", "webm", "avi"):
        f[kind] = video(out, kind)
    return f


if __name__ == "__main__":
    target = Path(sys.argv[1] if len(sys.argv) > 1 else "fixtures-out")
    for name, path in build_all(target).items():
        print(f"{name}\t{path.name}\t{path.stat().st_size}")
