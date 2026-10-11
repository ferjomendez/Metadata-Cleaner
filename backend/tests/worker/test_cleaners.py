"""Cleaning tests for every supported format and both modes.

Run inside the worker-test image (needs MAT2, ExifTool, FFmpeg, PIL, pikepdf).
Outputs are verified independently of the code under test: ExifTool and
ffprobe are run again on the output, the raw bytes are searched for the
planted markers, and content is compared with format libraries.
"""

import json
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest

from app.cleaners.base import CleanError
from app.process_file import process
from tests.fixtures import build_fixtures as fx

ALL_FORMATS = ["jpeg", "png", "webp", "tiff", "gif", "heic", "pdf", "docx", "xlsx", "pptx", "odt", "ods", "odp",
               "mp3", "wav", "flac", "m4a", "aac", "ogg", "mp4", "mov", "mkv", "webm", "avi"]


@pytest.fixture(scope="session")
def fixtures(tmp_path_factory):
    return fx.build_all(tmp_path_factory.mktemp("fixtures"))


def clean(settings, fixtures, fmt, mode, tmp_path, src=None):
    if fmt not in fixtures and src is None:
        pytest.skip(f"fixture {fmt} not available in this environment")
    src = src or fixtures[fmt]
    work = tmp_path / f"work-{fmt}-{mode}"
    work.mkdir()
    staged = tmp_path / f"in-{fmt}-{mode}"
    shutil.copyfile(src, staged)  # incoming files have no extension
    return process({"src": str(staged), "workdir": str(work), "format": fmt, "mode": mode})


def exif_json(path: Path) -> str:
    r = subprocess.run(["exiftool", "-a", "-G1", "-j", "-ee", str(path)], capture_output=True, check=False)
    return r.stdout.decode("utf-8", "replace")


def ffprobe_json(path: Path) -> str:
    r = subprocess.run(["ffprobe", "-v", "error", "-show_format", "-show_streams", "-print_format", "json",
                        str(path)], capture_output=True, check=False)
    return r.stdout.decode("utf-8", "replace")


def zip_text(path: Path) -> bytes:
    with zipfile.ZipFile(path) as zf:
        return b"\n".join(zf.read(n) for n in zf.namelist() if n.endswith((".xml", ".rels", ".vml")))


def check_codes(result):
    return {c["code"]: c["ok"] for c in result["report"]["checks"]}


@pytest.mark.parametrize("mode", ["conservative", "deep"])
@pytest.mark.parametrize("fmt", ALL_FORMATS)
def test_markers_removed_everywhere(settings, fixtures, tmp_path, fmt, mode):
    """Planted author/serial markers must not survive in any format/mode
    according to ExifTool (and ffprobe for media)."""
    r = clean(settings, fixtures, fmt, mode, tmp_path)
    assert r["ok"], r
    out = Path(r["output"])
    meta = exif_json(out)
    if fmt in ("mp3", "wav", "flac", "m4a", "aac", "ogg", "mp4", "mov", "mkv", "webm", "avi"):
        meta += ffprobe_json(out)
    for marker in fx.MARKERS:
        assert marker not in meta, f"{marker} still present in {fmt}/{mode}"
    assert "33.4489" not in meta and "GPSLatitude" not in meta
    assert all(c["ok"] for c in r["report"]["checks"]), r["report"]["checks"]
    assert r["report"]["removed"], "report lists removed items"
    assert r["report"]["not_verified"], "report lists not verified items"


# ---------------------------------------------------------------- images
@pytest.mark.parametrize("fmt", ["jpeg", "png", "webp", "tiff", "gif"])
def test_conservative_images_pixels_identical(settings, fixtures, tmp_path, fmt):
    from app.cleaners.images import frame_digests
    r = clean(settings, fixtures, fmt, "conservative", tmp_path)
    assert check_codes(r)["pixels_identical"]
    assert frame_digests(fixtures[fmt]) == frame_digests(Path(r["output"]))


def test_conservative_jpeg_keeps_orientation_and_icc(settings, fixtures, tmp_path):
    r = clean(settings, fixtures, "jpeg", "conservative", tmp_path)
    meta = json.loads(exif_json(Path(r["output"])))[0]
    assert meta.get("IFD0:Orientation") == "Rotate 90 CW"
    assert any(k.startswith("ICC") for k in meta), "ICC profile kept"
    assert "IFD1:ThumbnailImage" not in meta and "IFD1:ThumbnailOffset" not in meta


def test_gif_animation_preserved(settings, fixtures, tmp_path):
    from PIL import Image
    for mode in ("conservative", "deep"):
        r = clean(settings, fixtures, "gif", mode, tmp_path)
        with Image.open(r["output"]) as im:
            assert im.n_frames == 3


def test_deep_jpeg_reencoded_and_displayed_upright(settings, fixtures, tmp_path):
    """The fixture has Orientation=6. Either MAT2 rotates the pixels (size
    swapped, no tag) or the tag is restored on unrotated pixels."""
    from PIL import Image
    r = clean(settings, fixtures, "jpeg", "deep", tmp_path)
    assert "w_jpeg_reencoded" in r["report"]["warnings"]
    with Image.open(r["output"]) as a, Image.open(fixtures["jpeg"]) as b:
        orient = a.getexif().get(0x0112, 1)
        if a.size == b.size[::-1]:
            assert orient == 1
        else:
            assert a.size == b.size and orient == 6
            assert check_codes(r)["orientation_restored"]


def test_image_pixel_limit(settings, fixtures, tmp_path, monkeypatch):
    from PIL import Image
    big = tmp_path / "big.png"
    Image.new("L", (1200, 1000)).save(big)
    monkeypatch.setenv("MAX_IMAGE_MEGAPIXELS", "1")
    with pytest.raises(CleanError) as e:
        clean(settings, fixtures, "png", "conservative", tmp_path, src=big)
    assert e.value.code == "image_too_large"


# ---------------------------------------------------------------- pdf
def test_pdf_conservative_preserves_text_and_annotations(settings, fixtures, tmp_path):
    import pikepdf
    r = clean(settings, fixtures, "pdf", "conservative", tmp_path)
    codes = check_codes(r)
    assert codes["content_streams_identical"] and codes["document_metadata_absent"]
    with pikepdf.open(r["output"]) as p:
        assert "/Info" not in p.trailer and "/Metadata" not in p.Root
        assert len(p.pages) == 2
        content = p.pages[0].Contents.read_bytes()
        assert fx.BODY_TEXT.encode() in content
        annot = p.pages[0].Annots[0]
        assert str(annot.Contents) == fx.COMMENT_TEXT and "/T" not in annot and "/M" not in annot
        assert len(p.attachments) == 1  # retained and reported
    raw = Path(r["output"]).read_bytes()
    assert fx.AUTHOR.encode() not in raw and fx.SERIAL.encode() not in raw
    keys = {i["key"] for i in r["report"]["retained"]}
    assert {"PDF-Structure:Annotations", "PDF-Structure:JavaScript", "PDF-Structure:EmbeddedFiles"} <= keys


def test_pdf_deep_rasterizes(settings, fixtures, tmp_path):
    import pikepdf
    r = clean(settings, fixtures, "pdf", "deep", tmp_path)
    codes = check_codes(r)
    assert codes["page_count_preserved"] and codes["text_layer_removed"]
    assert "w_pdf_rasterized" in r["report"]["warnings"]
    with pikepdf.open(r["output"]) as p:
        assert len(p.pages) == 2 and not p.attachments
    raw = Path(r["output"]).read_bytes()
    assert fx.AUTHOR.encode() not in raw and fx.COMMENT_TEXT.encode() not in raw


def test_pdf_encrypted_rejected(settings, fixtures, tmp_path):
    with pytest.raises(CleanError) as e:
        clean(settings, fixtures, "pdf", "conservative", tmp_path, src=fixtures["pdf_encrypted"])
    assert e.value.code == "encrypted"


def test_pdf_corrupt(settings, fixtures, tmp_path):
    bad = tmp_path / "bad.pdf"
    bad.write_bytes(fixtures["pdf"].read_bytes()[:200])
    with pytest.raises(CleanError) as e:
        clean(settings, fixtures, "pdf", "conservative", tmp_path, src=bad)
    assert e.value.code == "corrupt"


# ---------------------------------------------------------------- office
def test_docx_conservative(settings, fixtures, tmp_path):
    import docx
    r = clean(settings, fixtures, "docx", "conservative", tmp_path)
    out = Path(r["output"])
    xml = zip_text(out)
    assert fx.AUTHOR.encode() not in xml and fx.SERIAL.encode() not in xml
    assert b"fer/Documentos" not in xml and b"rsidR" not in xml
    assert fx.COMMENT_TEXT.encode() in xml          # comment kept...
    assert b'w:author="Autor 1"' in xml             # ...author pseudonymised
    assert b"texto insertado" in xml                 # tracked change kept
    with zipfile.ZipFile(out) as zf:
        names = zf.namelist()
        assert "docProps/custom.xml" not in names
        assert {i.date_time for i in zf.infolist()} == {(1980, 1, 1, 0, 0, 0)}
        media = [n for n in names if n.startswith("word/media/")]
        img = tmp_path / "embedded.jpg"
        img.write_bytes(zf.read(media[0]))
    assert fx.AUTHOR not in exif_json(img) and "GPS" not in exif_json(img)
    d = docx.Document(str(out))
    assert any(fx.BODY_TEXT in p.text for p in d.paragraphs)
    assert check_codes(r)["content_preserved"] and check_codes(r)["package_valid"]


def test_docx_deep(settings, fixtures, tmp_path):
    import docx
    r = clean(settings, fixtures, "docx", "deep", tmp_path)
    out = Path(r["output"])
    xml = zip_text(out)
    assert fx.AUTHOR.encode() not in xml and fx.COMMENT_TEXT.encode() not in xml
    assert b"texto insertado" in xml  # insertion accepted into the text
    d = docx.Document(str(out))  # MAT2 alone produced a dangling relationship here
    assert any(fx.BODY_TEXT in p.text for p in d.paragraphs)


def test_xlsx_both_modes_keep_formulas(settings, fixtures, tmp_path):
    import openpyxl
    for mode in ("conservative", "deep"):
        r = clean(settings, fixtures, "xlsx", mode, tmp_path)
        out = Path(r["output"])
        xml = zip_text(out)
        assert fx.AUTHOR.encode() not in xml and b"fer\\Documentos" not in xml
        wb = openpyxl.load_workbook(out)
        assert wb["Datos"]["A3"].value == "=A1*A2" and wb["Datos"]["C1"].value == "=SUM(A1:A3)"
        assert wb["Oculta"].sheet_state == "hidden"
        keys = {i["key"] for i in r["report"]["retained"]}
        assert "Office-Structure:hidden_sheets" in keys
        if mode == "conservative":
            assert fx.COMMENT_TEXT.encode() in xml


def test_pptx_both_modes(settings, fixtures, tmp_path):
    from pptx import Presentation
    for mode in ("conservative", "deep"):
        r = clean(settings, fixtures, "pptx", mode, tmp_path)
        out = Path(r["output"])
        assert fx.AUTHOR.encode() not in zip_text(out)
        prs = Presentation(str(out))
        assert prs.slides[0].shapes.title.text == fx.BODY_TEXT
        assert len(prs.slides) == 2


@pytest.mark.parametrize("kind", ["odt", "ods", "odp"])
def test_odf_both_modes(settings, fixtures, tmp_path, kind):
    for mode in ("conservative", "deep"):
        r = clean(settings, fixtures, kind, mode, tmp_path)
        out = Path(r["output"])
        with zipfile.ZipFile(out) as zf:
            infos = zf.infolist()
            assert infos[0].filename == "mimetype" and infos[0].compress_type == zipfile.ZIP_STORED
            assert not any(n.startswith("Thumbnails/") for n in zf.namelist())
            content = zf.read("content.xml")
        assert fx.AUTHOR.encode() not in zip_text(out)
        assert fx.BODY_TEXT.encode() in content
        if mode == "conservative":
            assert fx.COMMENT_TEXT.encode() in content


def test_office_zip_bomb_rejected(settings, fixtures, tmp_path, monkeypatch):
    monkeypatch.setenv("MAX_ARCHIVE_RATIO", "10")
    bomb = tmp_path / "bomb.docx"
    shutil.copyfile(fixtures["docx"], bomb)
    with zipfile.ZipFile(bomb, "a", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("word/media/zeros.bin", b"\0" * (3 * 1024 * 1024))
    with pytest.raises(CleanError) as e:
        clean(settings, fixtures, "docx", "conservative", tmp_path, src=bomb)
    assert e.value.code == "archive_bomb"


# ---------------------------------------------------------------- media
@pytest.mark.parametrize("fmt", ["mp3", "wav", "flac", "m4a", "aac", "ogg"])
def test_audio_streams_identical(settings, fixtures, tmp_path, fmt):
    for mode in ("conservative", "deep"):
        r = clean(settings, fixtures, fmt, mode, tmp_path)
        assert check_codes(r)["streams_identical"]
        streams = json.loads(ffprobe_json(Path(r["output"])))["streams"]
        assert not any(s.get("disposition", {}).get("attached_pic") for s in streams)
        assert {s["codec_type"] for s in streams} == {"audio"}


@pytest.mark.parametrize("fmt", ["mp4", "mov", "mkv", "webm", "avi"])
def test_video_conservative_stream_copy(settings, fixtures, tmp_path, fmt):
    r = clean(settings, fixtures, fmt, "conservative", tmp_path)
    assert check_codes(r)["streams_identical"]
    probe = json.loads(ffprobe_json(Path(r["output"])))
    assert "location" not in json.dumps(probe.get("format", {}).get("tags", {}))
    if fmt == "mkv":
        assert any(s["codec_type"] == "subtitle" for s in probe["streams"])


def test_video_deep_removes_x264_sei(settings, fixtures, tmp_path):
    assert b"x264 - core" in fixtures["mp4"].read_bytes()
    c = clean(settings, fixtures, "mp4", "conservative", tmp_path)
    assert b"x264 - core" in Path(c["output"]).read_bytes()  # conservative copies the bitstream as is
    d = clean(settings, fixtures, "mp4", "deep", tmp_path)
    assert b"x264 - core" not in Path(d["output"]).read_bytes()
    assert check_codes(d)["video_packet_count_preserved"] and check_codes(d)["audio_streams_identical"]


def test_mkv_deep_drops_subtitles(settings, fixtures, tmp_path):
    r = clean(settings, fixtures, "mkv", "deep", tmp_path)
    probe = json.loads(ffprobe_json(Path(r["output"])))
    assert {s["codec_type"] for s in probe["streams"]} == {"video", "audio"}


def test_format_mismatch_rejected(settings, fixtures, tmp_path):
    with pytest.raises(CleanError) as e:
        clean(settings, fixtures, "png", "conservative", tmp_path, src=fixtures["jpeg"])
    assert e.value.code == "format_mismatch"
